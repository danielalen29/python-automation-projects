from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import pandas as pandas


def clean_csv_save_worksheets(
    csv_to_clean: Path,
    report_path: Path,
    csv_files_folder: Path,
) -> None:
    dataframe = pandas.read_csv(
        csv_to_clean,
        dtype={'Quantity': 'Int64'},
    )

    missing_rows = dataframe[
        dataframe['Quantity'].isna() | dataframe['Price'].isna()
    ].copy()

    clean_dataframe = dataframe.dropna(
        subset=['Quantity', 'Price']
    ).copy()
    clean_dataframe['Revenue'] = clean_dataframe['Quantity'] * clean_dataframe['Price']

    summary_dataframe = clean_dataframe.groupby(
        'Product', as_index=False,
    )[['Quantity', 'Revenue']].sum()
    summary_dataframe = summary_dataframe.sort_values(
        'Revenue',
        ascending=False,
        ignore_index=True,
    )

    missing_rows.to_csv(
        csv_files_folder / 'missing_values.csv',
        index=True,
        index_label='SourceIndex',
    )
    clean_dataframe.to_csv(
        csv_files_folder / 'cleaned_sales.csv',
        index=False,
    )
    summary_dataframe.to_csv(
        csv_files_folder / 'sales_summary.csv',
        index=False,
    )

    with pandas.ExcelWriter(report_path) as writer:
        clean_dataframe.to_excel(writer, sheet_name='Cleaned Sales', index=False)
        summary_dataframe.to_excel(writer, sheet_name='Summary', index=False)


def format_workbook(
    report_path: Path,
):
    workbook = load_workbook(report_path)
    clean_worksheet = workbook['Cleaned Sales']
    summary_worksheet = workbook['Summary']

    for worksheet in (clean_worksheet, summary_worksheet):
        for cell in worksheet[1]:
            cell.font = Font(bold=True)

    for column in ('D', 'F'):
        for cell in clean_worksheet[column][1:]:
            cell.number_format = '$0.00'

    for cell in summary_worksheet['C'][1:]:
        cell.number_format = '$0.00'

    for worksheet in (clean_worksheet, summary_worksheet):
        last_column = get_column_letter(worksheet.max_column)

        worksheet.auto_filter.ref = f'A1:{last_column}{worksheet.max_row}'
        worksheet.freeze_panes = 'A2'

    for worksheet in (clean_worksheet, summary_worksheet):
        for column in worksheet.columns:
            maximum_length = 0

            for cell in column:
                if cell.value is not None:
                    cell_length = len(str(cell.value))

                    if cell_length > maximum_length:
                        maximum_length = cell_length

            column_letter = column[0].column_letter
            worksheet.column_dimensions[column_letter].width = maximum_length + 3

    workbook.save(report_path)