from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import pandas as pandas
from pandas import DataFrame



def load_file(
    csv_to_clean_convert: Path,
) -> DataFrame:
    dataframe = pandas.read_csv(Path(csv_to_clean_convert))

    return dataframe


def remove_save_duplicates(
    dataframe: DataFrame,
) -> DataFrame:
    duplicate_rows = dataframe[dataframe.duplicated(keep=False)].copy()
    duplicate_rows['Quantity'] = pandas.to_numeric(duplicate_rows['Quantity'], errors='coerce').astype('Int64')
    duplicate_rows['Price'] = pandas.to_numeric(duplicate_rows['Price'], errors="coerce").astype('Float64')

    dataframe = dataframe.drop_duplicates()

    return duplicate_rows


def clean_columns(
    dataframe: DataFrame,
) -> DataFrame:
    for column in ('Product', 'Category', 'Country'):
        dataframe[column] = dataframe[column].str.strip().str.title()

    dataframe['Price'] = dataframe['Price'].str.replace('$', '', regex=False)
    dataframe['Price'] = pandas.to_numeric(dataframe['Price'], errors='coerce')
    dataframe['Quantity'] = pandas.to_numeric(dataframe['Quantity'], errors='coerce').astype('Int64')

    return dataframe


def create_save_sheets(
    dataframe: DataFrame,
    report_path: Path,
    duplicate_rows: DataFrame,
) -> None:
    review_dataframe = dataframe[dataframe['Product'].isna() | dataframe['Quantity'].isna() | dataframe['Price'].isna()].copy()

    clean_dataframe = dataframe.dropna(subset=['Product', 'Quantity', 'Price']).copy()
    clean_dataframe['Revenue'] = clean_dataframe['Quantity'] * clean_dataframe['Price']

    summary_dataframe = clean_dataframe.groupby('Product', as_index=False)[['Quantity', 'Revenue']].sum()
    summary_dataframe = summary_dataframe.sort_values('Revenue', ascending=False, ignore_index=True)
    
    with pandas.ExcelWriter(report_path) as writer:
        clean_dataframe.to_excel(writer, sheet_name='Cleaned Sales', index=False)
        summary_dataframe.to_excel(writer, sheet_name='Summary', index=False)
        review_dataframe.to_excel(writer, sheet_name='To Review', index=False)
        duplicate_rows.to_excel(writer, sheet_name='Duplicates', index=False)


def format_workbook(
    report_path: Path,
) -> None:
    workbook = load_workbook(report_path)

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        headers = {}
        for cell in worksheet[1]:
            headers[cell.value] = cell.column
        for header in ('Price', 'Revenue'):
            if header in headers:
                column_number = headers[header]
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_number).number_format = '$0.00'

        last_column = get_column_letter(worksheet.max_column)
        worksheet.auto_filter.ref = f'A1:{last_column}{worksheet.max_row}'
        worksheet.freeze_panes = 'A2'

        for column in worksheet.columns:
            maximum_length = 0

            for cell in column:
                if cell.value is not None:
                    value_length = len(str(cell.value))

                    if value_length > maximum_length:
                        maximum_length = value_length

            worksheet.column_dimensions[column[0].column_letter].width = maximum_length + 3

    workbook.save(report_path)