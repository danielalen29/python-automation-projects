from pathlib import Path
import csv

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from typing import TypedDict


class Repository(TypedDict):
    Name: str
    Language: str
    Stars: int
    URL: str


def get_repositories(
    user: str,
    per_page: int = 100,
    sort: str = 'updated',
) -> list[Repository]:
    url = f'https://api.github.com/users/{user}/repos'
    page = 1
    repositories = []

    while True:
        params = {
            'per_page': per_page,
            'page': page,
            'sort': sort,
        }

        response = requests.get(url, params=params, timeout=10)
        response.raise_for_status()

        data = response.json()

        if not data:
            break

        for repo in data:
            details: Repository = {
                'Name': repo['name'],
                'Language': repo['language'] or 'Unknown',
                'Stars': repo['stargazers_count'],
                'URL': repo['html_url'],
            }

            repositories.append(details)

        page += 1
    return repositories


def save_repositories_to_csv(
    repositories: list[Repository],
    file_path: Path,
) -> None:
    fieldnames = ['Name', 'Language', 'Stars', 'URL']

    with open(file_path, 'w', newline='', encoding='utf-8') as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(repositories)


def create_repositories_sheet(
    workbook: Workbook,
    repositories: list[Repository],
) -> Worksheet:
    repositories_worksheet = workbook.active
    repositories_worksheet.title = 'Repositories'
    repositories_worksheet.freeze_panes = 'A2'

    headers = ['Name', 'Language', 'Stars', 'URL']
    repositories_worksheet.append(headers)

    for cell in repositories_worksheet[1]:
        cell.font = Font(bold=True)

    for repository in repositories:
        row = [repository[header] for header in headers]
        repositories_worksheet.append(row)

    column_letter = get_column_letter(headers.index('URL') + 1)
    for row_number in range(2, repositories_worksheet.max_row + 1):
        cell = repositories_worksheet[f'{column_letter}{row_number}']
        cell.hyperlink = cell.value
        cell.style = 'Hyperlink'

    last_column = get_column_letter(repositories_worksheet.max_column)
    repositories_worksheet.auto_filter.ref = f'A1:{last_column}{repositories_worksheet.max_row}'

    return repositories_worksheet


def create_summary_sheet(
    workbook: Workbook,
    repositories: list[Repository],
) -> Worksheet:
    summary_worksheet = workbook.create_sheet('Summary')

    summary_worksheet.freeze_panes = 'A2'

    summary_headers = ['Metric', 'Value']
    summary_worksheet.append(summary_headers)

    for cell in summary_worksheet[1]:
        cell.font = Font(bold=True)

    total_stars = 0
    language_counts = {}

    for repository in repositories:
        total_stars += repository['Stars']
        language = repository['Language']

        if language in language_counts:
            language_counts[language] += 1
        else:
            language_counts[language] = 1

    if language_counts:
        most_used_language = max(language_counts, key=language_counts.get)
    else:
        most_used_language = 'N/A'

    summary_worksheet.append(['Total Repositories', len(repositories)])
    summary_worksheet.append(['Total Stars', total_stars])
    summary_worksheet.append(['Most Used Language', most_used_language])

    summary_worksheet.append([])

    language_headers = ['Language', 'Repository Count']
    summary_worksheet.append(language_headers)

    language_headers_row = summary_worksheet.max_row

    for cell in summary_worksheet[language_headers_row]:
        cell.font = Font(bold=True)

    for language, count in sorted(
        language_counts.items(),
        key=lambda item: item[1],
        reverse=True,
    ):
        summary_worksheet.append([language, count])

    return summary_worksheet


def set_column_widths(
    worksheets: list[Worksheet],
) -> None:
    for worksheet in worksheets:
        for column in worksheet.columns:
            max_length = 0

            for cell in column:
                if cell.value is not None:
                    cell_length = len(str(cell.value))

                    if cell_length > max_length:
                        max_length = cell_length

            column_letter = column[0].column_letter
            worksheet.column_dimensions[column_letter].width = max_length + 2


def save_repositories_to_excel(
    repositories: list[Repository],
    file_path: Path,
) -> None:
    workbook = Workbook()

    repositories_sheet = create_repositories_sheet(workbook, repositories)
    summary_sheet = create_summary_sheet(workbook, repositories)
    set_column_widths([repositories_sheet, summary_sheet])

    workbook.save(file_path)